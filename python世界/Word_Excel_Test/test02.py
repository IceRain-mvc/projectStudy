import xlwt
import xlrd
from docx import Document

path = "D:\\test\\test.docx"
document = Document(path)
sn_item = "Manu. Serial Number"
sn_sj_item = "40GE"
cs = "Vendor Name"
tt = "Transceiver Type"

# 原有的行数据
list_line_name = xlrd.open_workbook("network_import_template.xlsx").sheets()[0].row_values(0)

# 初始化列表 装载相应的字段
list_sn = []
list_sn_sj = []
list_cs = []
list_tt = []
# dict_sn = {}
i = 0

for paragraph in document.paragraphs:
    # 每行所有的文本
    line = paragraph.text
    # 字段中包含sn
    index_sn = line.find(sn_item)
    index_sn_sj = line.find(sn_sj_item)
    index_cs = line.find(cs)
    index_tt = line.find(tt)
    if index_sn != -1:
        # 去空格
        new_line = line.strip()
        # 截取
        find_sn = new_line[len(sn_item):len(new_line)]
        # 去除冒号
        new_find_sn = find_sn.strip()[1:len(find_sn)]
        # 添加到列表中
        list_sn.append(new_find_sn)
    if index_sn_sj != -1:
        # 去除左右空格 按照空格分割 取出第0个元素
        new_line_item = line.strip().split(" ")[0]
        list_sn_sj.append(new_line_item)
    if index_cs != -1:
        # 去左右空格
        new_line = line.strip()
        # 截取
        find_cs = new_line[len(sn_item):len(new_line)]
        # 去除冒号
        new_find_sn = find_cs.strip()[1:len(find_cs)]
        # 添加到列表中
        if i % 2 == 0:
            list_cs.append(new_find_sn)
        i += 1
    if index_tt != -1:
        # 去左右空格
        new_line = line.strip()
        # 截取
        find_tt = new_line[len(sn_item):len(new_line)]
        # 去除冒号
        new_find_sn = find_tt.strip()[1:len(find_tt)]
        print(new_find_sn)
        # 添加到列表中
        list_tt.append(new_find_sn)

# 向excel表中插入标题栏
workbook = xlwt.Workbook(encoding="ascii")
worksheet = workbook.add_sheet("sheet0")

i = 0
worksheet.row(0).set_style(xlwt.easyxf('font:height 250;'))
for item in list_line_name:
    worksheet.col(i).width = 256 * 20
    worksheet.write(0, i, item)
    i += 1

# 接下来是插入
# 首先是插入到sn
i = 1
for sn_item in list_sn:
    # 第一项 插入到 1,1
    # 第二项 插入到 2,1
    worksheet.write(i, 1, sn_item)
    i += 1

i = 1
for sn_sj_item in list_sn_sj:
    worksheet.write(i, 2, sn_sj_item)
    i += 1

# 以下是5 6 列
i = 1
for cs_item in list_cs:
    worksheet.write(i, 5, cs_item)
    i += 1

i = 1
for tt_item in list_tt:
    worksheet.write(i, 6, tt_item)
    i += 1


try:
    workbook.save("my excel.xls")
except PermissionError as result:
    print(result)
    print("请关闭文件在做操作")
    pass
