import xlwt
import xlrd
from docx import Document

path = "D:\\test\\test.docx"
document = Document(path)
sn = "Manu. Serial Number"

# 原有的行数据
list_line_name = xlrd.open_workbook("network_import_template.xlsx").sheets()[0].row_values(0)

# 初始化列表 装载相应的字段
list_sn = []
# dict_sn = {}

for paragraph in document.paragraphs:
    # 每行所有的文本
    line = paragraph.text
    # 字段中包含sn
    index = line.find(sn)
    if index != -1:
        # 去空格
        new_line = line.strip()
        # 截取
        find_sn = new_line[len(sn):len(new_line)]
        # 去除冒号
        new_find_sn = find_sn.strip()[1:len(find_sn)]
        # 添加到列表中
        list_sn.append(new_find_sn)
        # 插入excel 的sn列中
        # 1,1 1,2 1,6 1,7
# dict_sn = {"list_sn": list_sn}

# 向excel表中插入标题栏
workbook = xlwt.Workbook(encoding="ascii")
worksheet = workbook.add_sheet("sheet0")

i = 0
worksheet.row(0).set_style(xlwt.easyxf('font:height 250;'))
for item in list_line_name:
    worksheet.col(i).width = 256 * 20
    worksheet.write(0, i, item)
    i += 1

# 接下来是插入
# 首先是插入到sn
i = 1
for sn in list_sn:
    # 第一项 插入到 1,1
    # 第二项 插入到 2,1
    worksheet.write(i, 1, sn)
    i += 1
    pass
workbook.save("my excel.xls")

